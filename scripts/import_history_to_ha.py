#!/usr/bin/env python3
"""
Import the Davis history spreadsheet into Home Assistant as long-term
*external statistics*, so the decades of records show up in statistics-graph
cards alongside the live data.

Reads Weather_History_Log.xlsx (tabs: Years, Months, Days, Hours) and creates
one statistic per metric per granularity, e.g.:
    weather:year_temp_high      "Hist Year Temp High"      (deg F)
    weather:month_baro_avg      "Hist Month Avg Barometer" (inHg)
    weather:hour_temp           "Hist Hour Temp"           (deg F)
Blank cells are skipped. Re-running overwrites the same points (no duplicates).

SETUP
  1. pip install openpyxl websocket-client
  2. HA -> your user (bottom-left) -> Security -> Long-lived access tokens ->
     Create Token -> paste into TOKEN below.
  3. Check XLSX_PATH points to your spreadsheet, and set TZ_OFFSET to your UTC offset.
  4. Run:  python import_history_to_ha.py
"""

import sys
import json
from datetime import datetime
import openpyxl
from websocket import create_connection   # pip install websocket-client

# ===================== CONFIG =====================
HA_HOST   = "192.168.1.100"   # your Home Assistant IP
HA_PORT   = 8123
TOKEN     = "PASTE_YOUR_HA_LONG_LIVED_ACCESS_TOKEN_HERE"
XLSX_PATH = r"Weather_History_Log.xlsx"
TZ_OFFSET = "-07:00"   # America/Phoenix (no daylight saving)
# ==================================================

# header prefix -> (metric suffix, friendly, unit)
METRICS = [
    ("Temp High",         "temp_high",     "Temp High",         "°F"),
    ("Temp Low",          "temp_low",      "Temp Low",          "°F"),
    ("Temp (",            "temp",          "Temp",              "°F"),
    ("Wind High",         "wind_high",     "Wind High",         "mph"),
    ("Highest Rain Rate", "rain_rate_max", "Highest Rain Rate", "in/h"),
    ("Total Rain",        "rain_total",    "Total Rain",        "in"),
    ("Humidity High",     "hum_high",      "Humidity High",     "%"),
    ("Humidity Low",      "hum_low",       "Humidity Low",      "%"),
    ("Humidity (",        "humidity",      "Humidity",          "%"),
    ("Avg Barometer",     "baro_avg",      "Avg Barometer",     "inHg"),
    ("Barometer (",       "barometer",     "Barometer",         "inHg"),
]

# tab name -> (granularity prefix, friendly prefix)
TABS = {
    "Years":  ("year",  "Year"),
    "Months": ("month", "Month"),
    "Days":   ("day",   "Day"),
    "Hours":  ("hour",  "Hour"),
}


def metric_for(header):
    if not header:
        return None
    for prefix, suffix, friendly, unit in METRICS:
        if str(header).startswith(prefix):
            return suffix, friendly, unit
    return None


def period_start(tab, label):
    if isinstance(label, datetime):
        d = label
    else:
        s = str(label).strip()
        if tab == "Years":
            d = datetime(int(s[:4]), 1, 1)
        elif tab == "Months":
            d = datetime(int(s[:4]), int(s[5:7]), 1)
        elif tab == "Days":
            d = datetime(int(s[:4]), int(s[5:7]), int(s[8:10]))
        else:
            raise ValueError("bad label %r" % (label,))
    return d.strftime("%Y-%m-%dT%H:00:00") + TZ_OFFSET


def build():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    series = {}
    for tab, (gran, gran_name) in TABS.items():
        if tab not in wb.sheetnames:
            continue
        ws = wb[tab]
        headers = [c.value for c in ws[1]]
        colmap = {}
        for ci, h in enumerate(headers):
            if ci == 0:
                continue
            m = metric_for(h)
            if m:
                colmap[ci] = m
        for row in ws.iter_rows(min_row=2, values_only=True):
            label = row[0]
            if label in (None, ""):
                continue
            start = period_start(tab, label)
            for ci, (suffix, friendly, unit) in colmap.items():
                val = row[ci] if ci < len(row) else None
                if val in (None, ""):
                    continue
                try:
                    v = float(val)
                except (TypeError, ValueError):
                    continue
                sid = "weather:%s_%s" % (gran, suffix)
                if sid not in series:
                    series[sid] = {"name": "Hist %s %s" % (gran_name, friendly),
                                   "unit": unit, "stats": []}
                series[sid]["stats"].append({"start": start, "mean": v, "min": v, "max": v})
    for s in series.values():
        s["stats"].sort(key=lambda d: d["start"])
    return series


def main():
    if TOKEN.startswith("PASTE"):
        sys.exit("Set TOKEN to a Home Assistant long-lived access token first.")
    series = build()
    total = sum(len(s["stats"]) for s in series.values())
    print("%d statistics, %d data points to import." % (len(series), total))

    conn = create_connection("ws://%s:%d/api/websocket" % (HA_HOST, HA_PORT), timeout=15)
    try:
        conn.recv()
        conn.send(json.dumps({"type": "auth", "access_token": TOKEN}))
        if json.loads(conn.recv()).get("type") != "auth_ok":
            sys.exit("Auth failed - check the token.")
        print("Authenticated.")
        mid = 1
        for sid, s in sorted(series.items()):
            conn.send(json.dumps({
                "id": mid, "type": "recorder/import_statistics",
                "metadata": {"has_mean": True, "has_sum": False,
                             "statistic_id": sid, "unit_of_measurement": s["unit"],
                             "source": "weather", "name": s["name"]},
                "stats": s["stats"],
            }))
            resp = json.loads(conn.recv())
            print("  %s: %d pts -> %s" % (sid, len(s["stats"]), "OK" if resp.get("success") else resp))
            mid += 1
    finally:
        conn.close()
    print("Done.")


if __name__ == "__main__":
    main()
